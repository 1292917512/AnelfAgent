"""文档解析与 uploads 索引链路测试。"""

from __future__ import annotations

import asyncio
import threading
import time
from pathlib import Path

import pytest

from agent.memory.doc_extract import extract_document_text
from agent.memory.embedding import Embedder
from agent.memory.memory_store import MemoryStore
from agent.memory.memory_sync import index_single_file, sync_files
from agent.memory.memory_utils import list_indexable_files, pack_embedding


@pytest.fixture
def memory_dir(tmp_path) -> Path:
    ws = tmp_path / "custom-data"
    ws.mkdir(parents=True)
    (ws / "note.md").write_text("# 便签\n这是记忆文件内容", encoding="utf-8")
    return ws


@pytest.fixture
def uploads(tmp_path) -> Path:
    up = tmp_path / "uploads"
    docs = up / "docs"
    docs.mkdir(parents=True)
    (docs / "doc1.txt").write_text("上传文档一的内容，包含若干文字。", encoding="utf-8")
    return up


def _make_docx(path: Path, text: str) -> None:
    import docx
    document = docx.Document()
    document.add_paragraph(text)
    document.save(str(path))


class TestExtractDocumentText:
    def test_txt_and_md(self, tmp_path) -> None:
        p = tmp_path / "a.txt"
        p.write_text("纯文本内容", encoding="utf-8")
        assert extract_document_text(p) == "纯文本内容"

    def test_docx(self, tmp_path) -> None:
        p = tmp_path / "a.docx"
        _make_docx(p, "Word 文档段落内容")
        assert "Word 文档段落内容" in extract_document_text(p)

    def test_xlsx(self, tmp_path) -> None:
        """xlsx 逐表提取：表名标题 + 制表符单元格行，空表跳过。"""
        from openpyxl import Workbook

        wb = Workbook()
        ws1 = wb.active
        ws1.title = "成绩"
        ws1.append(["姓名", "分数"])
        ws1.append(["张三", 95])
        wb.create_sheet("空表")
        wb.save(str(tmp_path / "t.xlsx"))
        text = extract_document_text(tmp_path / "t.xlsx")
        assert "[成绩]" in text
        assert "姓名\t分数" in text
        assert "张三\t95" in text
        assert "空表" not in text  # 空表跳过

    def test_pptx(self, tmp_path) -> None:
        """pptx 逐页提取：文本框 + 备注，空页跳过。"""
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "季度汇报"
        slide.placeholders[1].text = "收入增长 20%"
        prs.save(str(tmp_path / "t.pptx"))
        text = extract_document_text(tmp_path / "t.pptx")
        assert "[幻灯片 1]" in text
        assert "季度汇报" in text
        assert "收入增长 20%" in text

    def test_empty_xlsx_raises(self, tmp_path) -> None:
        """全空工作簿无可用文本，显式报错。"""
        from openpyxl import Workbook

        wb = Workbook()
        wb.save(str(tmp_path / "empty.xlsx"))
        with pytest.raises(ValueError, match="无可用文本"):
            extract_document_text(tmp_path / "empty.xlsx")

    def test_unsupported_ext(self, tmp_path) -> None:
        p = tmp_path / "a.exe"
        p.write_bytes(b"fake")
        with pytest.raises(ValueError, match="不支持"):
            extract_document_text(p)

    def test_invalid_pdf_raises(self, tmp_path) -> None:
        p = tmp_path / "broken.pdf"
        p.write_bytes(b"not a real pdf")
        with pytest.raises(Exception):
            extract_document_text(p)


class TestListIndexableFiles:
    def test_namespaces(self, memory_dir: Path, uploads: Path) -> None:
        pairs = dict((rel, p) for p, rel in list_indexable_files(memory_dir, uploads))
        assert "memory/note.md" in pairs
        assert "uploads/docs/doc1.txt" in pairs

    def test_no_uploads_dir(self, memory_dir: Path) -> None:
        pairs = list_indexable_files(memory_dir, None)
        assert [rel for _, rel in pairs] == ["memory/note.md"]


async def _chunk_texts(store: MemoryStore, path: str) -> list[str]:
    db = await store._get_db()
    cursor = await db.execute("SELECT text FROM chunks WHERE path=?", (path,))
    return [row["text"] for row in await cursor.fetchall()]


class TestSyncFilesWithDocs:
    async def test_docs_indexed(self, store: MemoryStore, memory_dir: Path, uploads: Path) -> None:
        stats = await sync_files(store, Embedder(), memory_dir, uploads_dir=uploads)
        assert stats["synced"] == 2

        files = {f["path"] for f in await store.list_files()}
        assert files == {"memory/note.md", "uploads/docs/doc1.txt"}

        texts = await _chunk_texts(store, "uploads/docs/doc1.txt")
        assert any("上传文档一的内容" in t for t in texts)

    async def test_removed_doc_cleaned(self, store: MemoryStore, memory_dir: Path, uploads: Path) -> None:
        await sync_files(store, Embedder(), memory_dir, uploads_dir=uploads)
        (uploads / "docs" / "doc1.txt").unlink()

        stats = await sync_files(store, Embedder(), memory_dir, uploads_dir=uploads)
        assert stats["removed"] == 1
        files = {f["path"] for f in await store.list_files()}
        assert files == {"memory/note.md"}
        assert await _chunk_texts(store, "uploads/docs/doc1.txt") == []

    async def test_docx_indexed(self, store: MemoryStore, memory_dir: Path, uploads: Path) -> None:
        _make_docx(uploads / "docs" / "report.docx", "季度报告正文内容")
        stats = await sync_files(store, Embedder(), memory_dir, uploads_dir=uploads)
        assert stats["synced"] == 3
        texts = await _chunk_texts(store, "uploads/docs/report.docx")
        assert any("季度报告正文内容" in t for t in texts)

    async def test_incremental_skip_unchanged(self, store: MemoryStore, memory_dir: Path, uploads: Path) -> None:
        await sync_files(store, Embedder(), memory_dir, uploads_dir=uploads)
        stats = await sync_files(store, Embedder(), memory_dir, uploads_dir=uploads)
        assert stats["synced"] == 0
        assert stats["unchanged"] == 2
        assert stats["failed"] == 0

    async def test_office_formats_and_partial_failure(self, store, memory_dir, uploads) -> None:
        from openpyxl import Workbook
        from pptx import Presentation

        workbook = Workbook()
        workbook.active.append(["回款计划", 120])
        workbook.save(uploads / "docs" / "plan.xlsx")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "执行安排"
        presentation.save(uploads / "docs" / "plan.pptx")
        (uploads / "docs" / "broken.pdf").write_bytes(b"invalid pdf")
        stats = await sync_files(store, Embedder(), memory_dir, uploads_dir=uploads)
        assert stats["synced"] == 4
        assert stats["failed"] == 1
        assert "回款计划" in "".join(await _chunk_texts(store, "uploads/docs/plan.xlsx"))
        assert "执行安排" in "".join(await _chunk_texts(store, "uploads/docs/plan.pptx"))


@pytest.mark.parametrize("mode", ["fts", "like", "vector", "unified"])
async def test_file_retrieval_preserves_tail_evidence_and_lines(store, mode) -> None:
    text = "不相关的背景介绍。\n" * 100 + "重试预算 应当保留这条关键事实。\n" + "附录内容\n" * 40
    await store.upsert_chunks([{
        "id": "chunk", "path": "memory/retry.md", "start_line": 30,
        "end_line": 30 + text.count("\n"), "hash": "content", "text": text,
        "embedding": pack_embedding([1.0, 0.0]), "updated_ns": time.time_ns(),
    }])
    if mode == "like":
        store._files._conn.chunks_fts_available = False
    if mode == "vector":
        store._files._conn.vec_available = False
        results = await store.search_chunks_vector([1.0, 0.0], query="重试预算")
    elif mode == "unified":
        results = [r.model_dump() for r in await store.search_unified("重试预算", [1.0, 0.0])]
    else:
        results = await store.search_chunks_fts("重试预算")
    assert len(results) == 1
    result = results[0]
    assert "重试预算" in result["snippet"]
    assert len(result["snippet"]) <= 500
    assert 30 < result["start_line"] <= 130 <= result["end_line"]
    assert result["provenance"]["chunk_start_line"] == 30
    assert result["provenance"]["truncated"]


@pytest.mark.parametrize("cancel", [False, True])
async def test_failed_file_replacement_preserves_previous_index(store, memory_dir, monkeypatch, cancel) -> None:
    path = memory_dir / "note.md"
    await index_single_file(store, Embedder(), path, "memory/note.md")
    old_file = await store.get_file("memory/note.md")
    old_chunks = await _chunk_texts(store, "memory/note.md")
    write_chunks = store._files._write_chunks

    async def fail_after_write(*args):
        await write_chunks(*args)
        raise asyncio.CancelledError() if cancel else RuntimeError("write interrupted")

    monkeypatch.setattr(store._files, "_write_chunks", fail_after_write)
    path.write_text("更新后的内容", encoding="utf-8")
    with pytest.raises(asyncio.CancelledError if cancel else RuntimeError):
        await index_single_file(store, Embedder(), path, "memory/note.md")
    assert await store.get_file("memory/note.md") == old_file
    assert await _chunk_texts(store, "memory/note.md") == old_chunks
    assert await store.search_chunks_fts("记忆文件")
    assert not await store.search_chunks_fts("更新")


async def test_document_preparation_does_not_block_event_loop(store, memory_dir, monkeypatch) -> None:
    from agent.memory import memory_sync

    started, release = threading.Event(), threading.Event()
    read = memory_sync._read_indexable_content

    def wait_for_release(path):
        started.set()
        if not release.wait(2):
            raise TimeoutError("document preparation blocked the event loop")
        return read(path)

    monkeypatch.setattr(memory_sync, "_read_indexable_content", wait_for_release)
    task = asyncio.create_task(index_single_file(store, Embedder(), memory_dir / "note.md", "memory/note.md"))
    try:
        assert await asyncio.to_thread(started.wait, 1)
        assert not task.done()
    finally:
        release.set()
    assert await task > 0
