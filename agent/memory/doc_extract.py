"""文档文本提取：PDF / Word / Excel / PPT / 纯文本 → 纯文本，供记忆索引与文件读取使用。"""

from __future__ import annotations

from pathlib import Path

# 可索引的文档扩展名（.md 同时属于便签体系）
SUPPORTED_DOC_EXTS = frozenset({".pdf", ".docx", ".txt", ".md", ".xlsx", ".pptx"})


def extract_document_text(path: Path) -> str:
    """提取文档纯文本内容。不支持的类型或无可用文本时抛 ValueError。"""
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".xlsx":
        return _extract_xlsx(path)
    if ext == ".pptx":
        return _extract_pptx(path)
    raise ValueError(f"不支持的文档类型: {ext or path.name}")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    text = "\n\n".join(p for p in pages if p)
    if not text:
        raise ValueError(f"PDF 无可用文本: {path.name}")
    return text


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    if not text:
        raise ValueError(f"DOCX 无可用文本: {path.name}")
    return text


def _extract_xlsx(path: Path) -> str:
    """工作簿逐表提取：表名标题 + 制表符分隔的单元格行（空表跳过）。"""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        blocks: list[str] = []
        for ws in wb.worksheets:
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                if any(cells):
                    rows.append("\t".join(cells))
            if rows:
                blocks.append(f"[{ws.title}]\n" + "\n".join(rows))
        if not blocks:
            raise ValueError(f"XLSX 无可用文本: {path.name}")
        return "\n\n".join(blocks)
    finally:
        wb.close()


def _extract_pptx(path: Path) -> str:
    """演示文稿逐页提取：幻灯片文本框 + 备注（空页跳过）。"""
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text:
                    lines.append(text)
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                lines.append(f"[备注] {notes}")
        if lines:
            blocks.append(f"[幻灯片 {i}]\n" + "\n".join(lines))
    if not blocks:
        raise ValueError(f"PPTX 无可用文本: {path.name}")
    return "\n\n".join(blocks)
